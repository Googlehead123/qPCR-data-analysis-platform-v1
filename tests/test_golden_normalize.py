"""Tests for the golden-output normalizer, lock check and test-data root."""

import importlib.metadata
import io
import zipfile

import numpy as np
import openpyxl
import pandas as pd
import plotly.graph_objects as go
from PIL import Image

from tests.golden import env, normalize, paths


def test_lock_mismatches_reports_difference(monkeypatch):
    real_version = importlib.metadata.version
    monkeypatch.setattr(
        importlib.metadata,
        "version",
        lambda name: "0.0.1" if name == "numpy" else real_version(name),
    )
    assert "numpy: installed 0.0.1 != locked 2.5.2" in env.lock_mismatches()


def test_lock_mismatches_reports_missing_and_skips_foreign_markers(tmp_path):
    req = tmp_path / "requirements.txt"
    req.write_text(
        "# comment\n"
        "    # via something\n"
        "not-a-real-pkg-xyz==1.0\n"
        "win-only-pkg==1.0 ; sys_platform == 'win32' and python_version < '2'\n"
    )
    assert env.lock_mismatches(req) == [
        "not-a-real-pkg-xyz: installed <missing> != locked 1.0"
    ]


def test_test_data_root_prefers_env(monkeypatch, tmp_path):
    monkeypatch.setenv("QPCR_TEST_DATA", str(tmp_path))
    assert paths.test_data_root() == tmp_path
    assert paths.recipes_dir() == tmp_path / "golden" / "recipes"
    assert paths.local_out("abc123") == tmp_path / "golden" / "out" / "abc123"


def test_test_data_root_none_outside_repo(monkeypatch, tmp_path):
    monkeypatch.delenv("QPCR_TEST_DATA", raising=False)
    monkeypatch.chdir(tmp_path)
    assert paths.test_data_root() is None
    assert paths.recipes_dir() is None
    assert paths.local_out("abc123") is None


def _xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "2026-10-29 14:32"
    ws["B2"] = 0.1234567890123456
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_cells_one_line_each_and_dates_preserved():
    parts = normalize.normalize_xlsx(_xlsx_bytes())
    lines = parts["sheet:Data"].splitlines()
    assert "A1\t2026-10-29 14:32" in lines
    assert "B2\t0.123456789012" in lines
    core = parts["docProps/core.xml"]
    assert "<dcterms:created" in core
    assert core.count("<TS>") == 2
    assert "<TS></dcterms:created>" in core
    assert "<TS></dcterms:modified>" in core


def test_version_token_masked():
    assert normalize.mask_text("Software: qPCR Analysis Suite v3.1") == "Software: <APP_VERSION>"
    assert (
        normalize.mask_text("qPCR Analysis Suite 2026.10 (commit 1a2b3c4-dirty)")
        == "<APP_VERSION>"
    )


def _png(pixel: tuple) -> bytes:
    img = Image.new("RGBA", (10, 20), (255, 255, 255, 255))
    img.putpixel((3, 4), pixel)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _pptx(png: bytes) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("ppt/slides/slide1.xml", "<p:sld xmlns:p='urn:x'><p:a>hi</p:a></p:sld>")
        z.writestr("ppt/media/image1.png", png)
        z.writestr("[Content_Types].xml", "<Types/>")
    return buf.getvalue()


def test_pptx_media_hash_changes_with_pixels():
    a = normalize.normalize_pptx(_pptx(_png((0, 0, 0, 255))))
    b = normalize.normalize_pptx(_pptx(_png((1, 0, 0, 255))))
    ka = a["ppt/media/image1.png"]
    kb = b["ppt/media/image1.png"]
    assert ka.startswith("<image 10x20 sha256:")
    assert kb.startswith("<image 10x20 sha256:")
    assert ka != kb
    assert a["ppt/slides/slide1.xml"] == b["ppt/slides/slide1.xml"]
    assert "\n" in a["ppt/slides/slide1.xml"]  # pretty-printed


def test_figure_json_is_canonical():
    f1 = go.Figure(go.Bar(x=["a", "b"], y=[1.0, 2.0], name="s"))
    f2 = go.Figure(go.Bar(name="s", y=[1.0, 2.0], x=["a", "b"]))
    f3 = go.Figure(go.Bar(x=["a", "b"], y=[1.0, 2.5], name="s"))
    out1 = normalize.normalize_figures({"f": f1})
    assert out1 == normalize.normalize_figures({"f": f2})
    assert out1 != normalize.normalize_figures({"f": f3})


def test_figure_floats_rounded_to_12_digits():
    f1 = go.Figure(go.Bar(x=["a"], y=[0.1234567890123456]))
    f2 = go.Figure(go.Bar(x=["a"], y=[0.1234567890123459]))
    assert normalize.normalize_figures({"f": f1}) == normalize.normalize_figures({"f": f2})


def test_frames_csv_rounded():
    df = pd.DataFrame({"a": [0.1234567890123456], "b": ["x"]})
    assert normalize.normalize_frames({"t": df})["t"] == ",a,b\n0,0.123456789012,x\n"


def _np_fig(y: list[float]) -> go.Figure:
    return go.Figure(go.Bar(x=["a", "b"], y=np.array(y), error_y={"array": np.array(y)}))


def test_numpy_backed_figure_floats_rounded_and_decoded():
    base = normalize.normalize_figures({"f": _np_fig([0.1234567890123456, 2.0])})["f"]
    ulp = normalize.normalize_figures({"f": _np_fig([0.1234567890123459, 2.0])})["f"]
    real = normalize.normalize_figures({"f": _np_fig([0.1234567891, 2.0])})["f"]
    assert "bdata" not in base
    assert base == ulp
    assert base != real
