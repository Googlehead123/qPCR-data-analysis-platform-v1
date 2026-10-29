"""Every reported number must reproduce Min's hand sheet.

The hand sheet works on each well's fold change (Min, 2026-10-28):

    ΔCt_i  = Ct_target_i - mean housekeeping Ct of that condition
    ΔΔCt_i = ΔCt_i - mean ΔCt of the ΔΔCt reference condition
    FC_i   = 2^-ΔΔCt_i

and reports, per condition, the simple mean of FC_i, SD (n-1) and SEM = SD/√n
of FC_i, and a two-tailed Student's t-test (Excel T.TEST(control, sample, 2, 2))
on the FC_i of the two conditions compared. Until 2026-10-28 the bar was
2^-(mean ΔΔCt), SD/SEM were Ct-domain and the test ran on ΔCt; until 2026-10-23
the test was also Welch's, which made p-values higher than the hand calculation
(90/90 comparisons on the test_data plates).

Every expected value is derived by hand from the Ct table below — never by
calling the function under test or scipy.stats.ttest_ind. The pinned literals
are what the hand sheet produces for these wells.
"""
import numpy as np
import pandas as pd
import pytest
from scipy import special, stats

# Ct values: 3 wells per gene per condition. "None" is the negative control and
# the ΔΔCt reference, "B" a specified comparison condition. A's target wells are
# deliberately wider than the control's, where Welch and Student diverge.
CT = {
    #        HK (ACTIN)             target G
    "None": ([18.00, 18.10, 17.90], [25.00, 25.05, 24.95]),
    "A":    ([18.20, 18.00, 18.10], [23.00, 23.60, 22.70]),
    "B":    ([18.00, 18.00, 18.00], [24.00, 24.20, 24.10]),
}


def _fc(ref="None", hk_override=None):
    """Hand-sheet per-well fold changes for every condition."""
    hk_mean = {c: np.mean(hk) for c, (hk, _) in CT.items()}
    hk_mean.update(hk_override or {})
    dct = {c: np.array(g) - hk_mean[c] for c, (_, g) in CT.items()}
    ref_dct = dct[ref].mean()
    return {c: 2.0 ** -(d - ref_dct) for c, d in dct.items()}


FC = _fc()


def _student_p(ref, cond):
    """Pooled-variance two-tailed t-test, written out (Excel T.TEST type 2)."""
    ref, cond = np.asarray(ref, float), np.asarray(cond, float)
    n1, n2 = ref.size, cond.size
    sp2 = ((n1 - 1) * ref.var(ddof=1) + (n2 - 1) * cond.var(ddof=1)) / (n1 + n2 - 2)
    t = (ref.mean() - cond.mean()) / np.sqrt(sp2 * (1 / n1 + 1 / n2))
    return 2 * special.stdtr(n1 + n2 - 2, -abs(t))


def _welch_p(ref, cond):
    """Welch-Satterthwaite two-tailed t-test, written out (Excel T.TEST type 3)."""
    ref, cond = np.asarray(ref, float), np.asarray(cond, float)
    v1, v2 = ref.var(ddof=1) / ref.size, cond.var(ddof=1) / cond.size
    t = (ref.mean() - cond.mean()) / np.sqrt(v1 + v2)
    df = (v1 + v2) ** 2 / (v1 ** 2 / (ref.size - 1) + v2 ** 2 / (cond.size - 1))
    return 2 * special.stdtr(df, -abs(t))


def _raw():
    rows = []
    for cond, (hk, g) in CT.items():
        for i, ct in enumerate(hk, 1):
            rows.append((f"{cond}-HK{i}", cond, "ACTIN", ct))
        for i, ct in enumerate(g, 1):
            rows.append((f"{cond}-G{i}", cond, "G", ct))
    return pd.DataFrame(rows, columns=["Well", "Sample", "Target", "CT"])


MAPPING = {c: {"condition": c} for c in CT}


def _run(compare="None", excluded_wells=None, ddct_ref="None", **kw):
    from qpcr.analysis import AnalysisEngine
    raw = _raw()
    excluded_wells = excluded_wells or {}
    processed = AnalysisEngine.calculate_ddct(
        raw, "ACTIN", ddct_ref, excluded_wells, set(), MAPPING
    )
    stats_df = AnalysisEngine.calculate_statistics(
        processed, compare, raw_data=raw, hk_gene="ACTIN",
        sample_mapping=MAPPING, excluded_wells=excluded_wells, **kw,
    )
    return stats_df.set_index("Condition")


def _points(excluded_wells=None):
    from qpcr.analysis import AnalysisEngine
    return AnalysisEngine.compute_replicate_fold_changes(
        _raw(), "ACTIN", "None", MAPPING, excluded_wells or {}, set()
    )


class TestSummaryMatchesHandSheet:
    def test_bar_is_simple_mean_of_per_well_fold_change(self, mock_streamlit):
        row = _run().loc["A"]
        assert row["Fold_Change"] == pytest.approx(FC["A"].mean(), rel=1e-12)
        assert row["Fold_Change"] == pytest.approx(4.13118420599, rel=1e-9)
        assert row["Relative_Expression"] == row["Fold_Change"]

    def test_sd_and_sem_are_over_per_well_fold_change(self, mock_streamlit):
        row = _run().loc["A"]
        sd = FC["A"].std(ddof=1)                                  # STDEV, n-1
        assert row["SD"] == pytest.approx(sd, rel=1e-12)
        assert row["SD"] == pytest.approx(1.23222215405, rel=1e-9)
        assert row["SEM"] == pytest.approx(sd / np.sqrt(3), rel=1e-12)
        assert row["SEM"] == pytest.approx(0.711423792342, rel=1e-9)

    def test_error_bars_are_symmetric_in_fold_units(self, mock_streamlit):
        row = _run().loc["A"]
        sd = FC["A"].std(ddof=1)
        sem = sd / np.sqrt(3)
        assert row["FC_Error_Upper"] == row["FC_Error_Lower"] == pytest.approx(sd)
        assert row["FC_SEM_Upper"] == row["FC_SEM_Lower"] == pytest.approx(sem)
        ci = stats.t.ppf(0.975, 2) * sem
        assert row["FC_CI_Upper"] == row["FC_CI_Lower"] == pytest.approx(ci)

    def test_reference_mean_is_about_one_not_forced_to_one(self, mock_streamlit):
        """mean(2^-x_i) >= 2^-mean(x_i): the hand sheet's control is 1.0004."""
        row = _run().loc["None"]
        assert row["Fold_Change"] == pytest.approx(FC["None"].mean(), rel=1e-12)
        assert row["Fold_Change"] == pytest.approx(1.00040041759, rel=1e-9)

    def test_ct_columns_stay_ct_domain(self, mock_streamlit):
        row = _run().loc["A"]
        assert row["Target_Ct_Mean"] == pytest.approx(23.1)
        assert row["Target_Ct_SD"] == pytest.approx(np.sqrt(0.21))
        assert row["HK_Ct_Mean"] == pytest.approx(18.1)
        assert row["Delta_Ct"] == pytest.approx(5.0)
        assert row["Delta_Delta_Ct"] == pytest.approx(-2.0)

    def test_single_well_has_no_spread(self, mock_streamlit):
        excluded = {("G", "A"): {"A-G2", "A-G3"}}
        row = _run(excluded_wells=excluded).loc["A"]
        assert row["n_replicates"] == 1
        assert row["Fold_Change"] == pytest.approx(FC["A"][0])
        for col in ("SD", "SEM", "FC_Error_Upper", "FC_SEM_Upper", "FC_CI_Upper"):
            assert pd.isna(row[col]), col


class TestStudentTTestOnFoldChange:
    def test_default_vs_negative_control(self, mock_streamlit):
        row = _run().loc["A"]
        assert row["p_value"] == pytest.approx(_student_p(FC["None"], FC["A"]), rel=1e-9)
        assert row["p_value"] == pytest.approx(0.0117003895956, rel=1e-9)
        assert row["significance"] == "*"
        assert row["stat_test_used"] == "Student t-test (n=3 vs ref n=3)"

    def test_default_is_not_welch(self, mock_streamlit):
        """Welch gives 0.0478 on these wells — the 2026-10-23 regression."""
        row = _run().loc["A"]
        assert row["p_value"] < _welch_p(FC["None"], FC["A"]) / 3

    @pytest.mark.parametrize("ref", ["None", "B"])
    def test_any_reference_condition(self, mock_streamlit, ref):
        """'Condition vs None, or vs a specified condition' — same test either way."""
        row = _run(compare=ref).loc["A"]
        assert row["p_value"] == pytest.approx(_student_p(FC[ref], FC["A"]), rel=1e-9)

    def test_second_comparison(self, mock_streamlit):
        row = _run(compare_condition_2="B").loc["A"]
        assert row["p_value_2"] == pytest.approx(_student_p(FC["B"], FC["A"]), rel=1e-9)
        assert row["p_value_2"] == pytest.approx(0.0341083340838, rel=1e-9)
        assert row["stat_test_used_2"].startswith("Student t-test")

    def test_p_value_does_not_depend_on_the_ddct_reference(self, mock_streamlit):
        """Normalising to another reference scales every FC_i of the gene by one
        constant, which leaves the t statistic unchanged."""
        p_none = _run(ddct_ref="None").loc["A", "p_value"]
        p_b = _run(ddct_ref="B").loc["A", "p_value"]
        assert p_b == pytest.approx(p_none, rel=1e-9)

    def test_welch_remains_available(self, mock_streamlit):
        row = _run(ttest_type="welch").loc["A"]
        assert row["p_value"] == pytest.approx(_welch_p(FC["None"], FC["A"]), rel=1e-9)
        assert row["stat_test_used"].startswith("Welch t-test")

    def test_unknown_test_name_is_rejected(self, mock_streamlit):
        """Anything but 'student' used to fall through to Welch silently."""
        with pytest.raises(ValueError, match="ttest_type"):
            _run(ttest_type="studnet")

    def test_default_constant_is_student(self, mock_streamlit):
        from qpcr.analysis import DEFAULT_TTEST_TYPE, TTEST_TYPES
        assert DEFAULT_TTEST_TYPE == "student"
        assert set(TTEST_TYPES) == {"student", "welch"}


class TestQcReducedReplicates:
    def test_two_wells_after_qc(self, mock_streamlit):
        """QC drops A's 23.60 well: 2 wells vs 3, pooled df = 3."""
        excluded = {("G", "A"): {"A-G2"}}
        row = _run(excluded_wells=excluded).loc["A"]
        a2 = FC["A"][[0, 2]]
        assert row["n_replicates"] == 2
        assert row["Fold_Change"] == pytest.approx(a2.mean(), rel=1e-12)
        assert row["SD"] == pytest.approx(a2.std(ddof=1), rel=1e-12)
        assert row["SEM"] == pytest.approx(0.495468896473, rel=1e-9)
        assert row["Target_Ct_SD"] == pytest.approx(np.sqrt(0.045))   # Ct, (23.0, 22.7)
        assert row["p_value"] == pytest.approx(_student_p(FC["None"], a2), rel=1e-9)
        assert row["p_value"] == pytest.approx(0.00199890419997, rel=1e-9)
        assert row["stat_test_used"] == "Student t-test (n=2 vs ref n=3)"

    def test_excluded_hk_well_moves_bar_and_test_together(self, mock_streamlit):
        excluded = {("ACTIN", "A"): {"A-HK1"}}       # drop HK 18.20 -> mean 18.05
        row = _run(excluded_wells=excluded).loc["A"]
        fa = _fc(hk_override={"A": 18.05})["A"]
        assert row["HK_Ct_Mean"] == pytest.approx(18.05)
        assert row["Fold_Change"] == pytest.approx(fa.mean(), rel=1e-12)
        assert row["p_value"] == pytest.approx(_student_p(FC["None"], fa), rel=1e-9)


class TestExcelSummaryIsTheHandSheetTable:
    def test_summary_has_one_row_per_condition_with_hand_values(self, mock_streamlit):
        """The Summary sheet used to group condition-level fold changes by
        Target/Group and report their mean/SD/count and the MINIMUM p — a mean
        of means, not anything the hand sheet computes. It must instead list
        each condition's own n, mean FC, SD, SEM, p and marker."""
        import io
        import openpyxl
        from importlib import import_module
        export_to_excel = import_module("streamlit qpcr analysis v1").export_to_excel

        res = _run(compare_condition_2="B").reset_index()
        params = {"Housekeeping_Gene": "ACTIN", "Reference_Sample": "None"}
        wb = openpyxl.load_workbook(io.BytesIO(
            export_to_excel(_raw(), {"G": res}, params, MAPPING)
        ))
        ws = wb["Summary"]
        hdr = [c.value for c in ws[1]]
        assert hdr[:8] == ["Target", "Condition", "n", "Fold_Change", "SD",
                           "SEM", "p_value", "significance"]
        assert "p_value_2" in hdr and "significance_2" in hdr
        rows = {r[1]: dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)}
        assert set(rows) == set(CT)
        a = rows["A"]
        assert a["n"] == 3
        assert a["Fold_Change"] == pytest.approx(FC["A"].mean(), rel=1e-9)
        assert a["SD"] == pytest.approx(FC["A"].std(ddof=1), rel=1e-9)
        assert a["SEM"] == pytest.approx(FC["A"].std(ddof=1) / np.sqrt(3), rel=1e-9)
        assert a["p_value"] == pytest.approx(_student_p(FC["None"], FC["A"]), rel=1e-9)
        assert a["significance"] == "*"
        assert a["p_value_2"] == pytest.approx(_student_p(FC["B"], FC["A"]), rel=1e-9)


class TestEverythingUsesTheSameWells:
    """Bar, error bar, scatter points (Replicate_FC) and p-value must all be
    computed from one set of per-well fold changes."""

    @pytest.mark.parametrize("excluded", [{}, {("G", "A"): {"A-G2"}},
                                          {("ACTIN", "A"): {"A-HK1"}}])
    def test_points_reproduce_bar_sd_and_p(self, mock_streamlit, excluded):
        res = _run(excluded_wells=excluded)
        pts = _points(excluded)
        by_cond = {c: g["Replicate_FC"].to_numpy() for c, g in pts.groupby("Condition")}
        for cond in CT:
            row = res.loc[cond]
            v = by_cond[cond]
            assert len(v) == row["n_replicates"]
            assert row["Fold_Change"] == pytest.approx(v.mean(), rel=1e-12)
            assert row["SD"] == pytest.approx(v.std(ddof=1), rel=1e-12)
        assert res.loc["A", "p_value"] == pytest.approx(
            _student_p(by_cond["None"], by_cond["A"]), rel=1e-9
        )
