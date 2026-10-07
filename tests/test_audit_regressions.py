"""Regression tests for the 2026-10 audit (Claude subagents + Codex).

Each test pins one defect that was reproduced against the code before the fix.
"""
import io

import numpy as np
import pandas as pd

from qpcr.analysis import AnalysisEngine
from qpcr.graph import GraphGenerator
from qpcr.parser import QPCRParser
from qpcr.quality_control import QualityControl


def _csv(rows, header="Well Position,Sample Name,Target Name,CT"):
    return io.BytesIO((header + "\n" + "\n".join(rows) + "\n").encode("utf-8"))


class TestGrubbsKeyedPerGroup:
    def test_outlier_flag_does_not_spill_to_other_samples_sharing_a_well(self):
        data = pd.DataFrame(
            {
                "Well": ["A1", "A2", "A3", "A1", "A2", "A3"],
                "Sample": ["S1"] * 3 + ["S2"] * 3,
                "Target": ["G"] * 6,
                "CT": [20.0, 20.1, 28.0, 20.0, 20.0, 20.0],
            }
        )
        qc = QualityControl.detect_outliers(data)
        flagged = qc[qc["Issues"].str.contains("Grubbs")]
        assert set(zip(flagged["Sample"], flagged["Well"])) == {("S1", "A3")}


class TestThresholdsForwarded:
    def test_summary_stats_health_uses_custom_ct_high(self):
        wells = [f"A{i}" for i in range(1, 4)]
        data = pd.DataFrame(
            {"Well": wells, "Sample": ["S"] * 3, "Target": ["G"] * 3,
             "CT": [36.0, 36.1, 35.9]}
        )
        stats = QualityControl.get_qc_summary_stats(
            data, thresholds={"ct_high": 40}
        )
        assert stats["high_ct_count"] == 0
        assert stats["healthy_triplicates"] == stats["total_triplicates"]


class TestHousekeepingCaseConsistency:
    def test_lowercase_hk_name_still_normalises(self):
        rows = []
        for cond, ct in (("Ctl", 25.0), ("Trt", 24.0)):
            rows += [
                {"Well": f"A{i}", "Sample": cond, "Condition": cond,
                 "Target": "GAPDH", "CT": 18.0}
                for i in range(1, 4)
            ]
            rows += [
                {"Well": f"B{i}", "Sample": cond, "Condition": cond,
                 "Target": "G1", "CT": ct}
                for i in range(1, 4)
            ]
        data = pd.DataFrame(rows)
        mapping = {c: {"condition": c, "group": "g", "include": True}
                   for c in ("Ctl", "Trt")}
        res = AnalysisEngine.calculate_ddct(
            data, "gapdh", "Ctl", set(), set(), mapping
        )
        assert not res.empty
        trt = res[res["Condition"] == "Trt"].iloc[0]
        assert np.isclose(trt["Relative_Expression"], 2.0)


class TestParserRobustness:
    def test_sample_named_NA_is_kept(self):
        f = _csv(["A1,NA,G,20.0", "A2,NA,G,20.1", "A3,S2,G,21.0"])
        out = QPCRParser.parse(f)
        assert out is not None
        assert "NA" in set(out["Sample"])

    def test_undetermined_ct_is_still_dropped(self):
        f = _csv(["A1,S1,G,Undetermined", "A2,S1,G,20.1"])
        out = QPCRParser.parse(f)
        assert len(out) == 1

    def test_infinite_ct_is_dropped_not_turned_into_a_fold_change(self):
        f = _csv(["A1,S1,G,inf", "A2,S1,G,20.1", "A3,S1,G,20.2"])
        out = QPCRParser.parse(f)
        assert np.isfinite(out["CT"]).all()
        assert len(out) == 2

    def test_lowercase_headers_parse(self):
        f = _csv(["A1,S1,G,20.0"], header="well,sample,target,ct")
        out = QPCRParser.parse(f)
        assert out is not None and len(out) == 1


class TestRefLineWithoutErrorBars:
    def test_reference_line_with_error_bars_off_builds(self, processed_gene_data):
        fig = GraphGenerator.create_gene_graph(
            data=processed_gene_data,
            gene="COL1A1",
            settings={"show_error": False},
            ref_line_value=1.0,
        )
        assert len(fig.data) >= 1


class TestExcelWritesNamesVerbatim:
    def test_export_does_not_turn_uploaded_names_into_formulas(self, mock_streamlit):
        import openpyxl
        from importlib import import_module

        spec = import_module("streamlit qpcr analysis v1")
        processed = {
            "G1": pd.DataFrame({
                "Target": ["G1", "G1"],
                "Condition": ["Non-treated", "=1+1"],
                "Group": ["Negative Control", "Treatment"],
                "Fold_Change": [1.0, 2.0],
                "Relative_Expression": [1.0, 2.0],
                "SEM": [0.05, 0.1],
            })
        }
        raw = pd.DataFrame(
            [{"Well": "A1", "Sample": s_, "Target": "G1", "CT": 20.0}
             for s_ in ("Non-treated", "=1+1")]
        )
        mapping = {"Non-treated": {"condition": "Non-treated", "group": "Negative Control"},
                   "=1+1": {"condition": "=1+1", "group": "Treatment"}}
        xlsx = spec.export_to_excel(
            raw, processed,
            {"Housekeeping_Gene": "GAPDH", "Efficacy_Type": "Anti-Aging"}, mapping,
        )
        wb = openpyxl.load_workbook(io.BytesIO(xlsx.getvalue() if hasattr(xlsx, "getvalue") else xlsx))
        formulas = [
            (ws.title, c.coordinate, c.value)
            for ws in wb.worksheets for row in ws.iter_rows() for c in row
            if c.data_type == "f"
        ]
        assert formulas == []


class TestCodexReviewFollowups:
    def test_well_header_matched_case_insensitively_when_not_first(self):
        f = _csv(["S1,G,20.0,A1", "S1,G,20.1,A2"], header="sample,target,ct,well")
        out = QPCRParser.parse(f)
        assert list(out["Well"]) == ["A1", "A2"]

    def test_hk_exclusion_applies_to_case_variant_rows(self):
        rows = []
        for cond in ("Ctl", "Trt"):
            for i, ct in enumerate((18.0, 18.0, 30.0), start=1):
                rows.append({"Well": f"A{i}", "Sample": cond, "Condition": cond,
                             "Target": "gapdh", "CT": ct})
            for i in range(1, 4):
                rows.append({"Well": f"B{i}", "Sample": cond, "Condition": cond,
                             "Target": "G1", "CT": 25.0})
        data = pd.DataFrame(rows)
        mapping = {c: {"condition": c, "group": "g", "include": True}
                   for c in ("Ctl", "Trt")}
        excl = {("gapdh", c): {"A3"} for c in ("Ctl", "Trt")}
        res = AnalysisEngine.calculate_ddct(
            data, "GAPDH", "Ctl", excl, set(), mapping
        )
        trt = res[res["Condition"] == "Trt"].iloc[0]
        assert np.isclose(trt["Relative_Expression"], 1.0)


class TestVerdictDirectionVsComparisonControl:
    def test_ratio_is_taken_against_the_comparison_condition(self):
        from qpcr.auto.interpret import fold_vs_comparison

        gd = pd.DataFrame({"Condition": ["Ref", "Ctl", "Trt"],
                           "Fold_Change": [1.0, 4.0, 2.0]})
        rows = gd[gd["Condition"] == "Trt"]
        out = fold_vs_comparison(rows, gd, "Ctl", "Fold_Change")
        assert np.isclose(out.iloc[0], 0.5)  # below its own control: not "up"

    def test_falls_back_to_reference_relative_when_comparison_missing(self):
        from qpcr.auto.interpret import fold_vs_comparison

        gd = pd.DataFrame({"Condition": ["Ref", "Trt"], "Fold_Change": [1.0, 2.0]})
        rows = gd[gd["Condition"] == "Trt"]
        assert np.isclose(fold_vs_comparison(rows, gd, "Nope", "Fold_Change").iloc[0], 2.0)
        assert np.isclose(fold_vs_comparison(rows, gd, None, "Fold_Change").iloc[0], 2.0)
        assert np.isclose(fold_vs_comparison(rows, gd, "ref", "Fold_Change").iloc[0], 2.0)


class TestDuplicateWellIdentity:
    def test_repeated_wells_become_distinct_and_exclusion_hits_only_one(self):
        from qpcr.utils import make_well_ids_unique

        data = pd.DataFrame(
            {"Well": ["A1", "A2", "A1"], "Sample": ["S"] * 3,
             "Target": ["G"] * 3, "CT": [20.0, 20.0, 30.0]}
        )
        out, n = make_well_ids_unique(data)
        assert n == 1
        assert list(out["Well"]) == ["A1", "A2", "A1 (2)"]
        excl, audit = QualityControl.auto_select_replicates(out)
        assert excl == {("G", "S"): {"A1 (2)"}}
        assert set(audit[0]["kept_wells"]) == {"A1", "A2"}

    def test_unique_ids_are_left_alone(self):
        from qpcr.utils import make_well_ids_unique

        data = pd.DataFrame({"Well": ["A1", "A1"], "Sample": ["S1", "S2"],
                             "Target": ["G", "G"], "CT": [20.0, 21.0]})
        out, n = make_well_ids_unique(data)
        assert n == 0 and list(out["Well"]) == ["A1", "A1"]


class TestAutoQcTrimIsDeclared:
    def test_summary_states_trim_and_bias(self, mock_streamlit):
        from importlib import import_module

        spec = import_module("streamlit qpcr analysis v1")
        audit = [{"status": "trimmed", "dropped_wells": ["A3"]},
                 {"status": "unresolved", "dropped_wells": ["B3"]}]
        txt = spec.auto_qc_trim_summary(audit, 0.3)
        assert "1 group(s) trimmed" in txt and "1 still above" in txt
        assert "2 well(s) dropped" in txt and "anti-conservative" in txt

    def test_summary_when_nothing_trimmed(self, mock_streamlit):
        from importlib import import_module

        spec = import_module("streamlit qpcr analysis v1")
        assert "nothing trimmed" in spec.auto_qc_trim_summary([], 0.3)

    def test_provenance_and_miqe_carry_it(self, mock_streamlit):
        from importlib import import_module
        from qpcr.auto import build_miqe_checklist

        spec = import_module("streamlit qpcr analysis v1")
        prov = spec.build_provenance(
            efficacy="x", hk_gene="GAPDH", ref_condition="Ctl", cmp_conditions=[],
            ttest_type="welch", excluded_wells={}, excluded_samples=set(),
            n_genes=1, n_samples=2, timestamp="t",
            auto_qc_audit=[{"status": "trimmed", "dropped_wells": ["A3"]}],
            auto_qc_threshold=0.3,
        )
        assert "trimmed" in prov["auto_qc_trim"]
        assert "Auto-QC replicate trim" in spec.format_provenance_text(prov)
        assert "Automatic replicate trim" in build_miqe_checklist(prov)


def _frame(gene, ref_fc=1.0, trt_fc=2.0):
    return pd.DataFrame({
        "Target": [gene, gene],
        "Condition": ["Non-treated", "Treatment"],
        "Group": ["Negative Control", "Treatment"],
        "Fold_Change": [ref_fc, trt_fc],
        "Relative_Expression": [ref_fc, trt_fc],
        "SEM": [0.05, 0.1],
    })


def _export(spec, processed, **kw):
    genes = list(processed)
    raw = pd.DataFrame(
        [{"Well": "A1", "Sample": s_, "Target": g, "CT": 20.0}
         for g in genes for s_ in ("Non-treated", "Treatment")]
    )
    mapping = {"Non-treated": {"condition": "Non-treated", "group": "Negative Control"},
               "Treatment": {"condition": "Treatment", "group": "Treatment"}}
    out = spec.export_to_excel(
        raw, processed, {"Housekeeping_Gene": "GAPDH", "Efficacy_Type": "Anti-Aging"},
        mapping, **kw,
    )
    return out.getvalue() if hasattr(out, "getvalue") else out


class TestExcelExportFidelity:
    def test_two_genes_with_one_display_name_both_survive_in_fc_matrix(self, mock_streamlit):
        from importlib import import_module

        import openpyxl

        spec = import_module("streamlit qpcr analysis v1")
        xlsx = _export(
            spec, {"G1": _frame("G1", trt_fc=2.0), "G2": _frame("G2", trt_fc=5.0)},
            gene_display_names={"G1": "Same", "G2": "Same"},
        )
        wb = openpyxl.load_workbook(io.BytesIO(xlsx))
        ws = wb["FC_Matrix"]
        labels = [r[0].value for r in ws.iter_rows(min_row=2) if r[0].value]
        assert len(labels) == 2 and len(set(labels)) == 2

    def test_per_gene_axis_settings_reach_the_chart(self, mock_streamlit):
        import re
        import zipfile
        from importlib import import_module

        spec = import_module("streamlit qpcr analysis v1")
        xlsx = _export(
            spec, {"G1": _frame("G1")},
            graph_settings={"G1_y_min": 0.5, "G1_y_max": 5, "G1_y_log": True},
        )
        zf = zipfile.ZipFile(io.BytesIO(xlsx))
        chart = next(n for n in zf.namelist() if re.match(r"xl/charts/chart\d+\.xml", n))
        xml = zf.read(chart).decode("utf-8")
        assert re.search(r'<(?:c:)?max val="5(\.0)?"', xml)
        assert re.search(r'<(?:c:)?min val="0\.5"', xml)
        assert "logBase" in xml


class TestCodexReviewFollowups2:
    def test_exact_comparison_name_wins_over_case_variant(self):
        from qpcr.auto.interpret import fold_vs_comparison

        gd = pd.DataFrame({"Condition": ["CTL", "Ctl", "Trt"],
                           "Fold_Change": [1.0, 4.0, 2.0]})
        rows = gd[gd["Condition"] == "Trt"]
        assert np.isclose(fold_vs_comparison(rows, gd, "Ctl", "Fold_Change").iloc[0], 0.5)

    def test_infinite_comparison_falls_back(self):
        from qpcr.auto.interpret import fold_vs_comparison

        gd = pd.DataFrame({"Condition": ["Ctl", "Trt"], "Fold_Change": [np.inf, 2.0]})
        rows = gd[gd["Condition"] == "Trt"]
        assert np.isclose(fold_vs_comparison(rows, gd, "Ctl", "Fold_Change").iloc[0], 2.0)

    def test_generated_well_ids_never_collide_with_existing_ones(self):
        from qpcr.utils import make_well_ids_unique

        data = pd.DataFrame({"Well": ["A1", "A1", "A1 (2)"], "Sample": ["S"] * 3,
                             "Target": ["G"] * 3, "CT": [20.0, 20.1, 20.2]})
        out, n = make_well_ids_unique(data)
        assert out["Well"].is_unique and n == 1
        assert out["Well"].iloc[0] == "A1" and out["Well"].iloc[2] == "A1 (2)"

    def test_missing_wells_get_distinct_ids(self):
        from qpcr.utils import make_well_ids_unique

        data = pd.DataFrame({"Well": [pd.NA, pd.NA, "A1"], "Sample": ["S"] * 3,
                             "Target": ["G"] * 3, "CT": [20.0, 20.1, 20.2]})
        out, _ = make_well_ids_unique(data)
        assert out["Well"].is_unique and out["Well"].notna().all()

    def test_display_name_suffix_cannot_collide_with_another_label(self, mock_streamlit):
        from importlib import import_module

        spec = import_module("streamlit qpcr analysis v1")
        names = spec._unique_display_names(
            {"G1": 0, "G2": 0, "G3": 0},
            {"G1": "Same", "G2": "Same", "G3": "Same (G1)"},
        )
        assert len(set(names.values())) == 3
